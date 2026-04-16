#!/usr/bin/env python3

from __future__ import annotations

import json
import re
import sys
import urllib.request
from collections import Counter
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))
from ward_utils import WORKBOOK_TO_STATE, normalize_seed_lga  # noqa: E402


ROOT = Path("/Users/arinzeogbonna/Work/personal/spending")
SEED_DIR = ROOT / "packages/database/seed"
ELECTORAL_PAGE = "https://www.inecnigeria.org/electoral-constituencies-of-the-federation/"

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))
from ward_utils import normalize_name


def fetch(url: str, dest: Path) -> None:
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0",
        },
    )
    with urllib.request.urlopen(req) as response:  # noqa: S310
        dest.write_bytes(response.read())


def pick_sd_sheet(workbook):
    for name in workbook.sheetnames:
        upper = name.upper()
        if upper in {"SD", "SENATORIAL", "IMO SD"}:
            return workbook[name]
        if upper.endswith("SD") or "SENATORIAL" in upper:
            return workbook[name]
    return workbook[workbook.sheetnames[0]]


def extract_official_lga_counts(path: Path) -> dict[str, int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = pick_sd_sheet(workbook)

    lga_idx = None
    ra_idx = None
    counts: dict[str, int] = {}

    for row in sheet.iter_rows(values_only=True):
        values = [str(v).strip().upper() if isinstance(v, str) else v for v in row]

        if lga_idx is None or ra_idx is None:
            for idx, value in enumerate(values):
                if isinstance(value, str) and "LGA COMPOSITION" in value:
                    lga_idx = idx
                if isinstance(value, str) and (
                    "NO OF RAS" in value
                    or "NO OF RA" in value
                    or "NO. OF RAS" in value
                    or "NO. OF RA" in value
                ):
                    ra_idx = idx
            if lga_idx is not None and ra_idx is not None:
                continue

        if lga_idx is None or ra_idx is None or len(row) <= max(lga_idx, ra_idx):
            continue

        lga = row[lga_idx]
        ras = row[ra_idx]
        if not isinstance(lga, str) or not isinstance(ras, (int, float)):
            continue

        name = lga.strip()
        if name.upper() == "TOTAL":
            continue

        counts[normalize_name(name)] = int(ras)

    return counts


def load_seed_counts() -> tuple[dict[str, int], dict[str, dict[str, int]]]:
    wards = json.loads((SEED_DIR / "wards.json").read_text())
    by_state = Counter()
    by_lga: dict[str, Counter] = {}

    for ward in wards:
        lga_code = ward["lga_code"]
        if lga_code.startswith("akwa_ibom_"):
            state = "akwa_ibom"
        elif lga_code.startswith("cross_river_"):
            state = "cross_river"
        else:
            state = lga_code.split("_")[0]

        by_state[state] += 1
        by_lga.setdefault(state, Counter())[normalize_seed_lga(state, lga_code)] += 1

    return dict(by_state), {state: dict(counter) for state, counter in by_lga.items()}


def build_report(cache_dir: Path) -> tuple[dict, dict[str, dict[str, int]]]:
    electoral_page = cache_dir / "electoral-constituencies.html"
    fetch(ELECTORAL_PAGE, electoral_page)

    attachment_urls = re.findall(
        r"https://wp1\.inecnigeria\.org/electoral-constituencies-of-the-federation/[^\"<> ]+",
        electoral_page.read_text(errors="ignore"),
    )

    attachments_dir = cache_dir / "attachments"
    attachments_dir.mkdir(exist_ok=True)
    workbooks_dir = cache_dir / "workbooks"
    workbooks_dir.mkdir(exist_ok=True)

    workbook_urls: list[str] = []
    for url in attachment_urls:
        slug = url.rstrip("/").split("/")[-1]
        attachment_path = attachments_dir / f"{slug}.html"
        fetch(url, attachment_path)

        match = re.search(
            r"href='([^']+\.(?:xlsx|xls|csv))'",
            attachment_path.read_text(errors="ignore"),
            re.IGNORECASE,
        )
        if not match:
            continue
        workbook_url = match.group(1)
        if workbook_url.startswith("/"):
            workbook_url = f"https://wp1.inecnigeria.org{workbook_url}"
        workbook_urls.append(workbook_url)

    official_by_state: dict[str, int] = {}
    official_by_lga: dict[str, dict[str, int]] = {}

    for url in workbook_urls:
        filename = url.split("/")[-1]
        path = workbooks_dir / filename
        fetch(url, path)

        stem = path.stem
        state = WORKBOOK_TO_STATE.get(stem.upper()) or WORKBOOK_TO_STATE.get(stem)
        lga_counts = extract_official_lga_counts(path)
        official_by_lga[state] = lga_counts
        official_by_state[state] = sum(lga_counts.values())

    seed_by_state, seed_by_lga = load_seed_counts()

    report = {
        "source_page": ELECTORAL_PAGE,
        "seed_total": sum(seed_by_state.values()),
        "official_total": sum(official_by_state.values()),
        "delta": sum(seed_by_state.values()) - sum(official_by_state.values()),
        "states": [],
    }

    for state in sorted(official_by_state):
        seed_total = seed_by_state.get(state, 0)
        official_total = official_by_state[state]

        seed_lgas = seed_by_lga.get(state, {})
        official_lgas = official_by_lga[state]

        mismatches = []
        for lga_name in sorted(set(seed_lgas) | set(official_lgas)):
            seed_count = seed_lgas.get(lga_name)
            official_count = official_lgas.get(lga_name)
            if seed_count != official_count:
                mismatches.append(
                    {
                        "lga": lga_name,
                        "seed": seed_count,
                        "official": official_count,
                        "delta": None
                        if seed_count is None or official_count is None
                        else seed_count - official_count,
                    }
                )

        report["states"].append(
            {
                "state": state,
                "seed_total": seed_total,
                "official_total": official_total,
                "delta": seed_total - official_total,
                "mismatched_lgas": mismatches,
            }
        )

    return report, official_by_lga


def main() -> int:
    output_path = SEED_DIR / "inec-ward-ra-mismatch-report.json"
    with TemporaryDirectory(prefix="inec-ward-audit-") as temp_dir:
        report, official_counts = build_report(Path(temp_dir))

    output_path.write_text(json.dumps(report, indent=2) + "\n")
    official_path = SEED_DIR / "inec-ward-ra-official-counts.json"
    official_path.write_text(
        json.dumps({"source_page": ELECTORAL_PAGE, "counts": official_counts}, indent=2)
        + "\n"
    )

    print(f"Wrote {output_path}")
    print(
        f"seed_total={report['seed_total']} official_total={report['official_total']} delta={report['delta']}"
    )
    worst = sorted(
        report["states"],
        key=lambda item: abs(item["delta"]),
        reverse=True,
    )[:10]
    for item in worst:
        print(
            f"{item['state']}: seed={item['seed_total']} official={item['official_total']} delta={item['delta']}"
        )

    return 0


if __name__ == "__main__":
    sys.exit(main())
